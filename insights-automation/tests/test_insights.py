"""Tests:  python -m pytest tests -q   (ca. 30 Sekunden)"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from insights.beispieldaten import erzeuge, schreibe_beispiel  # noqa: E402
from insights.kennzahlen import (anteile_zeitreihe, beitragsanalyse, markt_ueberblick, ranking,  # noqa: E402
                                 segment_matrix, zeitraum)
from insights.kommentar import erkenntnisse, glaetten_mit_ollama, pp, zahl  # noqa: E402
from insights.laden import lade, pruefe  # noqa: E402
from insights.prognose import Szenario, prognostiziere  # noqa: E402


@pytest.fixture(scope="session")
def df() -> pd.DataFrame:
    return erzeuge()


# ------------------------------------------------------------------ Daten


def test_beispieldaten_deterministisch_und_vollstaendig(df):
    assert df.equals(erzeuge())
    assert df["monat"].nunique() == 48
    assert set(df["kategorie"]) == {"Zigaretten", "Tabakerhitzer", "E-Zigaretten"}
    assert (df["absatz"] > 0).all() and (df["umsatz"] > 0).all()
    assert df[df["marke"] == "Spark"]["monat"].min() == "2025-06"


def test_laden_ueber_mapping_ergibt_dieselben_daten(df, tmp_path):
    pfade = schreibe_beispiel(tmp_path)
    geladen, meta = lade(pfade["mapping"])
    assert meta["eigener_hersteller"] == "Eigen AG" and len(meta["quellen"]) == 2
    schluessel = ["monat", "region", "kanal", "marke"]
    a = df.sort_values(schluessel).reset_index(drop=True)
    b = geladen.sort_values(schluessel).reset_index(drop=True)
    for spalte in ("absatz", "umsatz", "distribution_gew", "werbedruck"):
        np.testing.assert_allclose(a[spalte], b[spalte], rtol=1e-9)


def test_laden_internes_format_csv_und_excel(df, tmp_path):
    df.to_csv(tmp_path / "d.csv", index=False)
    df.head(500).to_excel(tmp_path / "d.xlsx", index=False)
    assert len(lade(tmp_path / "d.csv")[0]) == len(df)
    assert len(lade(tmp_path / "d.xlsx")[0]) == 500


def test_pruefung_findet_eingebaute_fehler(df):
    kaputt = df.copy()
    kaputt = pd.concat([kaputt, kaputt.iloc[[0]]])                              # Dublette
    kaputt = kaputt[kaputt["monat"] != "2024-02"]                              # fehlender Monat
    kaputt.loc[kaputt.index[5], "absatz"] = -10                                # negativer Wert
    i = kaputt[(kaputt["marke"] == "Atlas") & (kaputt["monat"] == "2025-05")].index
    kaputt.loc[i, "umsatz"] *= 3                                               # Preissprung
    themen = {(p.stufe, p.thema) for p in pruefe(kaputt)}
    assert ("Fehler", "Dubletten") in themen
    assert ("Fehler", "Perioden") in themen
    assert ("Fehler", "Negative Werte") in themen
    assert ("Warnung", "Preissprung") in themen
    assert not [p for p in pruefe(df) if p.stufe in ("Fehler", "Warnung")]


# ------------------------------------------------------------------ Kennzahlen


def test_anteile_summieren_sich_auf_100(df):
    a = anteile_zeitreihe(df)
    summen = a.groupby(["kategorie", "monat"])["anteil"].sum()
    np.testing.assert_allclose(summen, 100)


def test_zeitraum_logik(df):
    z = zeitraum(df["monat"].unique(), "2026-03", 3)
    assert z.monate == ["2026-01", "2026-02", "2026-03"]
    assert z.vj == ["2025-01", "2025-02", "2025-03"]
    assert z.vp == ["2025-10", "2025-11", "2025-12"]


def test_ranking_konsistent(df):
    z = zeitraum(df["monat"].unique(), "2026-09", 3)
    r = ranking(df, "Zigaretten", z)
    assert r["anteil_menge"].sum() == pytest.approx(100)
    assert r["delta_menge_vj"].sum() == pytest.approx(0, abs=1e-9)  # Anteile verschieben sich nur
    # Preisindex: absatzgewichtet ergibt der Markt 100
    assert (r["preisindex"] * r["absatz"]).sum() / r["absatz"].sum() == pytest.approx(100, rel=0.02)


def test_beitragsanalyse_summiert_exakt(df):
    for ende in ("2025-06", "2026-09"):
        z = zeitraum(df["monat"].unique(), ende, 3)
        for kat, marke in (("Zigaretten", "Atlas"), ("Tabakerhitzer", "Ember")):
            national = ranking(df, kat, z).set_index("marke").loc[marke, "delta_menge_vj"]
            assert beitragsanalyse(df, kat, z, marke)["beitrag"].sum() == pytest.approx(national, abs=1e-9)


def test_beitragsanalyse_findet_listungsverlust(df):
    z = zeitraum(df["monat"].unique(), "2025-06", 3)
    b = beitragsanalyse(df, "Zigaretten", z, "Atlas")
    schlimmstes = b.iloc[0]
    assert (schlimmstes["region"], schlimmstes["kanal"]) == ("Sued", "Kiosk & Fachhandel")
    m = segment_matrix(df, "Zigaretten", z, "Atlas")
    assert m.set_index(["region", "kanal"]).loc[("Sued", "Kiosk & Fachhandel"), "delta_vj"] < -3


def test_ueberblick_aus_sicht_des_eigenen_herstellers(df):
    z = zeitraum(df["monat"].unique(), "2026-09", 3)
    u = markt_ueberblick(df, z, "Eigen AG").set_index("kategorie")
    assert u.loc["Zigaretten", "rang"] == 1
    assert u.loc["Tabakerhitzer", "markt_veraenderung_vj"] > 10  # wachsende Kategorie
    assert u.loc["Zigaretten", "markt_veraenderung_vj"] < 0      # schrumpfende Kategorie


# ------------------------------------------------------------------ Prognose


def test_prognose_mit_rueckblick_und_band(df):
    p = prognostiziere(df, "Zigaretten", "Atlas", horizont=12)
    t = p.tabelle
    assert len(t) == 12 and t["monat"].iloc[0] == "2026-10"
    assert (t["unten"] <= t["prognose"]).all() and (t["prognose"] <= t["oben"]).all()
    assert t["oben"].iloc[-1] - t["unten"].iloc[-1] > t["oben"].iloc[0] - t["unten"].iloc[0]  # Band waechst
    assert len(p.rueckblick) == 6 and p.rueckblick["mae_kombi"].iloc[0] < 1.5
    assert 0 <= p.gewicht_treiber <= 1


def test_szenarien_haben_plausible_richtung(df):
    p = prognostiziere(df, "Zigaretten", "Atlas", szenarien=[
        Szenario("teurer", preis_prozent=5), Szenario("mehr Listung", distribution_pp=3)])
    basis = p.tabelle["prognose"].iloc[-1]
    assert p.szenarien["teurer"]["prognose"].iloc[-1] < basis
    assert p.szenarien["mehr Listung"]["prognose"].iloc[-1] > basis
    assert p.sensitivitaet["Preis +1 %"] < 0 < p.sensitivitaet["Distribution +1 pp"]


def test_prognose_trifft_zukunft_im_rueckblick(df):
    """Prognose ab 2025-09 gegen die tatsaechlichen 6 Folgemonate (ohne Ereignisse dazwischen)."""
    bis = df[df["monat"] <= "2025-09"]
    p = prognostiziere(bis, "Zigaretten", "Eigen AG", horizont=6, ebene="hersteller")
    ist = (df[(df["kategorie"] == "Zigaretten") & (df["hersteller"] == "Eigen AG")].groupby("monat")["absatz"].sum()
           / df[df["kategorie"] == "Zigaretten"].groupby("monat")["absatz"].sum() * 100)
    fehler = (p.tabelle.set_index("monat")["prognose"] - ist.reindex(p.tabelle["monat"])).abs()
    assert fehler.mean() < 1.5


def test_kurze_historie_faellt_auf_trend_zurueck(df):
    p = prognostiziere(df, "E-Zigaretten", "Spark")
    assert p.gewicht_treiber == 0 and p.hinweis
    assert p.tabelle["unten"].notna().all()


# ------------------------------------------------------------------ Kommentare


def test_zahlenformat_deutsch():
    assert zahl(1234.56) == "1.234,6"
    assert pp(0.42) == "+0,4 pp" and pp(-1.25) == "−1,2 pp" or pp(-1.25) == "−1,3 pp"


def test_kommentare_finden_die_eingebauten_ereignisse(df):
    d = df[df["monat"] <= "2025-06"]
    texte = " ".join(e.text for e in erkenntnisse(d, zeitraum(d["monat"].unique(), "2025-06", 3)))
    assert "Atlas verliert" in texte and "Kiosk & Fachhandel Süd" in texte and "Distribution" in texte
    assert "Neue Marke Spark" in texte
    d = df[df["monat"] <= "2026-09"]
    erk = erkenntnisse(d, zeitraum(d["monat"].unique(), "2026-09", 3))
    ember = next(e.text for e in erk if e.text.startswith("Ember"))
    assert "gewinnt" in ember and "Werbedruck" in ember
    cirrus = [e.text for e in erk if e.text.startswith("Cirrus")]
    assert all("Werbedruck" not in t for t in cirrus)  # keine unpassenden Treiber bei Verlust


def test_ollama_glaetten_faellt_sicher_zurueck():
    texte = ["Atlas verliert 0,6 pp."]
    assert glaetten_mit_ollama(texte, url="http://127.0.0.1:9") == texte


# ------------------------------------------------------------------ Gesamtbericht


def test_bericht_ende_zu_ende(tmp_path):
    from openpyxl import load_workbook
    from pptx import Presentation

    from insights.bericht import erstelle

    pfade = schreibe_beispiel(tmp_path / "daten")
    erg = erstelle(pfade["mapping"], tmp_path / "out", hinweis="Test")
    prs = Presentation(erg.folien)
    assert len(prs.slides) == 11
    diagramme = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
    assert diagramme >= 7  # echte PowerPoint-Diagramme, keine Bilder
    titel = [next(sh.text_frame.text for sh in s.shapes if sh.name == "Titel") for s in prs.slides]
    assert all(t and not t.endswith(".") for t in titel)
    wb = load_workbook(erg.excel)
    assert {"Übersicht", "Erkenntnisse", "Zeitreihen", "Prognose", "Datenqualität"} <= set(wb.sheetnames)
    assert len(wb["Zeitreihen"]._charts) == 3
    erstes = wb["Zeitreihen"].cell(2, 1).value
    assert erstes == "monat" and wb["Zeitreihen"].cell(2, 9).value == "monat"  # Tabellen nebeneinander


def test_bericht_bricht_bei_datenfehlern_ab(df, tmp_path):
    from insights.bericht import erstelle

    kaputt = pd.concat([df, df.iloc[[0]]])
    kaputt.to_csv(tmp_path / "kaputt.csv", index=False)
    with pytest.raises(SystemExit, match="Datenfehler"):
        erstelle(tmp_path / "kaputt.csv", tmp_path / "out")

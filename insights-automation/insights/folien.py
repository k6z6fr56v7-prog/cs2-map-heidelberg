"""PowerPoint-Deck mit echten, bearbeitbaren PowerPoint-Diagrammen.

Aufbau (Pyramide): Kernaussagen zuerst, dann je Kategorie eine Folie mit
Aussage-Titel, Diagramm und Kommentar, dann Treiber (Segmente), Prognose,
Methodik. Sprechernotizen enthalten die Details fuer die Praesentation.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import config
from .kommentar import Erkenntnis, pp, pro_kategorie, zahl

SCHRIFT = "Calibri"
BREITE, HOEHE = Inches(13.333), Inches(7.5)
RAND = Inches(0.6)


def rgb(hexwert: str) -> RGBColor:
    return RGBColor.from_string(hexwert)


def _text(folie, x, y, w, h, text, groesse=16, fett=False, farbe=config.TINTE, ausrichtung=PP_ALIGN.LEFT,
          anker=MSO_ANCHOR.TOP, name=None):
    box = folie.shapes.add_textbox(x, y, w, h)
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anker
    tf.margin_left = tf.margin_right = Emu(0)
    absaetze = text if isinstance(text, list) else [text]
    for i, a in enumerate(absaetze):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ausrichtung
        lauf = p.add_run()
        lauf.text = a
        lauf.font.size, lauf.font.bold, lauf.font.name = Pt(groesse), fett, SCHRIFT
        lauf.font.color.rgb = rgb(farbe)
    return box


def _aufzaehlung(folie, x, y, w, h, punkte: list[str], groesse=15, farbe=config.TINTE):
    box = folie.shapes.add_textbox(x, y, w, h)
    box.name = "Kommentar"
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    for i, punkt in enumerate(punkte):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        # Echter Aufzaehlungspunkt ueber XML (python-pptx hat dafuer keine API)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.22))))
        pPr.set("indent", str(-Emu(Inches(0.22))))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        pPr.append(bu)
        lauf = p.add_run()
        lauf.text = punkt
        lauf.font.size, lauf.font.name = Pt(groesse), SCHRIFT
        lauf.font.color.rgb = rgb(farbe)
    return box


def _titel(folie, text, untertitel=None):
    _text(folie, RAND, Inches(0.45), BREITE - 2 * RAND, Inches(0.9), text, groesse=28, fett=True,
          farbe=config.DUNKEL, anker=MSO_ANCHOR.TOP, name="Titel")
    if untertitel:
        _text(folie, RAND, Inches(1.25), BREITE - 2 * RAND, Inches(0.4), untertitel, groesse=13,
              farbe=config.TINTE_2, name="Untertitel")


def _fusszeile(folie, text, seite):
    _text(folie, RAND, HOEHE - Inches(0.45), Inches(9), Inches(0.3), text, groesse=10,
          farbe=config.TINTE_GEDAEMPFT, name="Fusszeile")
    _text(folie, BREITE - RAND - Inches(1), HOEHE - Inches(0.45), Inches(1), Inches(0.3), str(seite),
          groesse=10, farbe=config.TINTE_GEDAEMPFT, ausrichtung=PP_ALIGN.RIGHT, name="Seitenzahl")


def _hintergrund(folie, farbe):
    fill = folie.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(farbe)


def _karte(folie, x, y, w, h, farbe="F3F6FB"):
    form = folie.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    form.adjustments[0] = 0.06
    form.fill.solid()
    form.fill.fore_color.rgb = rgb(farbe)
    form.line.fill.background()
    form.shadow.inherit = False
    return form


def _diagramm_stil(chart, legende=True):
    chart.font.name, chart.font.size = SCHRIFT, Pt(11)
    chart.font.color.rgb = rgb(config.TINTE_2)
    chart.has_legend = legende
    if legende:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)


def _achsen(chart, zahlformat='0"%"', min_=None, max_=None):
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(config.GITTER)
    va.format.line.fill.background()
    va.tick_labels.number_format = zahlformat
    va.tick_labels.number_format_is_linked = False
    if min_ is not None:
        va.minimum_scale = min_
    if max_ is not None:
        va.maximum_scale = max_
    ca = chart.category_axis
    ca.format.line.color.rgb = rgb("C3C2B7")
    ca.tick_labels.font.size = Pt(10)


# ---------------------------------------------------------------- Folien


def _titelfolie(prs, titel, untertitel, hinweis):
    f = prs.slides.add_slide(prs.slide_layouts[6])
    _hintergrund(f, config.DUNKEL)
    _text(f, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.6), titel, groesse=42, fett=True, farbe="FFFFFF",
          anker=MSO_ANCHOR.BOTTOM, name="Titel")
    _text(f, Inches(0.9), Inches(4.05), Inches(11.5), Inches(0.6), untertitel, groesse=20,
          farbe=config.HELL_AUF_DUNKEL, name="Untertitel")
    _text(f, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.4), hinweis, groesse=12,
          farbe=config.HELL_AUF_DUNKEL, name="Hinweis")
    return f


def _zusammenfassung(prs, ueberblick: pd.DataFrame, kernaussagen: list[str], fuss, seite):
    f = prs.slides.add_slide(prs.slide_layouts[6])
    _titel(f, "Auf einen Blick", f"Mengenanteile {config.EIGENER_HERSTELLER}, Veränderung vs. Vorjahr")
    n = len(ueberblick)
    luecke = Inches(0.35)
    w = (BREITE - 2 * RAND - luecke * (n - 1)) / n
    for i, (_, u) in enumerate(ueberblick.iterrows()):
        x = RAND + i * (w + luecke)
        y = Inches(1.9)
        _karte(f, x, y, w, Inches(2.0))
        _text(f, x + Inches(0.3), y + Inches(0.2), w - Inches(0.6), Inches(0.4), u["kategorie"], groesse=15,
              fett=True, farbe=config.TINTE_2)
        _text(f, x + Inches(0.3), y + Inches(0.55), w - Inches(0.6), Inches(0.9), f"{zahl(u['anteil_menge'])} %",
              groesse=44, fett=True, farbe=config.DUNKEL)
        d = u["delta_menge_vj"]
        pfeil = "▲" if d >= 0 else "▼"
        farbe = "006300" if d >= 0 else "B3261E"
        _text(f, x + Inches(0.3), y + Inches(1.45), w - Inches(0.6), Inches(0.4),
              f"{pfeil} {pp(d)} vs. VJ  ·  Rang {int(u['rang'])}  ·  Markt {u['markt_veraenderung_vj']:+.1f} %".replace(".", ","),
              groesse=13, farbe=farbe)
    _text(f, RAND, Inches(4.25), Inches(6), Inches(0.4), "Kernaussagen", groesse=18, fett=True, farbe=config.DUNKEL)
    _aufzaehlung(f, RAND, Inches(4.7), BREITE - 2 * RAND, Inches(2.2), kernaussagen, groesse=15)
    _fusszeile(f, fuss, seite)
    f.notes_slide.notes_text_frame.text = "Kernaussagen im Detail:\n" + "\n".join(kernaussagen)
    return f


def _kategorie_folie(prs, kat: str, zeitreihe: pd.DataFrame, eigene: list[str], punkte: list[Erkenntnis],
                     fuss, seite):
    f = prs.slides.add_slide(prs.slide_layouts[6])
    marke_punkte = [p for p in punkte if p.typ == "Marke"]
    if marke_punkte:
        kern = marke_punkte[0].text.split(" (jetzt")[0]
        titel = f"{kat}: {kern}"
    else:
        titel = f"{kat}: Anteile stabil"
    _titel(f, titel, "Mengenanteile in %, letzte 24 Monate, national")
    # Reihen: eigene Marken, die 2 groessten Wettbewerber, Rest zusammengefasst
    letzte = zeitreihe.tail(24)
    groesse = letzte.iloc[-1].sort_values(ascending=False)
    wettbewerb = [m for m in groesse.index if m not in eigene][:2]
    rest = [m for m in letzte.columns if m not in eigene and m not in wettbewerb]
    cd = CategoryChartData()
    cd.categories = [m[2:].replace("-", "/") for m in letzte.index]  # z.B. 24/09
    reihen = []
    for m in eigene:
        if m in letzte.columns:
            cd.add_series(m, [round(v, 2) if pd.notna(v) else None for v in letzte[m]])
            reihen.append((m, config.FARBE_EIGEN, 2.75))
    farben = iter(config.FARBEN_WETTBEWERB)
    for m in wettbewerb:
        cd.add_series(m, [round(v, 2) if pd.notna(v) else None for v in letzte[m]])
        reihen.append((m, next(farben), 2.0))
    if rest:
        cd.add_series("Übrige", [round(v, 2) for v in letzte[rest].fillna(0).sum(axis=1)])
        reihen.append(("Übrige", config.FARBE_UEBRIGE, 1.5))
    gf = f.shapes.add_chart(XL_CHART_TYPE.LINE, RAND, Inches(1.8), Inches(7.6), Inches(5.0), cd)
    gf.name = "Anteilsverlauf"
    chart = gf.chart
    _diagramm_stil(chart)
    _achsen(chart)
    eigen_strich = [None, "dash", "sysDot"]
    i_eigen = 0
    for serie, (name, farbe, breite) in zip(chart.series, reihen):
        serie.smooth = False
        serie.format.line.color.rgb = rgb(farbe)
        serie.format.line.width = Pt(breite)
        serie.marker.style = XL_MARKER_STYLE.NONE
        if farbe == config.FARBE_EIGEN:
            if eigen_strich[i_eigen % 3]:
                serie.format.line.dash_style = {"dash": MSO_LINE_DASH_STYLE.DASH,
                                                "sysDot": MSO_LINE_DASH_STYLE.ROUND_DOT}[eigen_strich[i_eigen % 3]]
            i_eigen += 1
    texte = [p.text for p in punkte if p.typ != "Ueberblick"][:4]
    if not texte:
        texte = [p.text for p in punkte][:3]
    _karte(f, Inches(8.55), Inches(1.8), BREITE - Inches(8.55) - RAND, Inches(5.0), "F7F7F5")
    _text(f, Inches(8.85), Inches(2.0), Inches(3.8), Inches(0.4), "Was passiert ist", groesse=16, fett=True,
          farbe=config.DUNKEL)
    _aufzaehlung(f, Inches(8.85), Inches(2.5), BREITE - Inches(8.85) - RAND - Inches(0.3), Inches(4.2),
                 [kuerze(t) for t in texte], groesse=13)
    _fusszeile(f, fuss, seite)
    f.notes_slide.notes_text_frame.text = "\n".join(p.text for p in punkte)
    return f


def kuerze(text: str, max_len: int = 230) -> str:
    return text if len(text) <= max_len else text[: max_len - 1].rsplit(" ", 1)[0] + " …"


def _gewinner_verlierer(prs, rankings: dict[str, pd.DataFrame], fuss, seite):
    f = prs.slides.add_slide(prs.slide_layouts[6])
    zeilen = []
    for kat, r in rankings.items():
        for _, z in r.iterrows():
            eigen = z["hersteller"] == config.EIGENER_HERSTELLER
            zeilen.append((f"{'★ ' if eigen else ''}{z['marke']} ({kat[:3]}.)", z["delta_menge_vj"], eigen))
    zeilen.sort(key=lambda t: t[1])
    if len(zeilen) > 14:  # nur die staerksten Bewegungen
        zeilen = sorted(sorted(zeilen, key=lambda t: -abs(t[1]))[:14], key=lambda t: t[1])
    gewinner = max(zeilen, key=lambda t: t[1])[0].split(" (")[0].lstrip("★ ")
    verlierer = min(zeilen, key=lambda t: t[1])[0].split(" (")[0].lstrip("★ ")
    _titel(f, f"Stärkster Gewinner {gewinner}, stärkster Verlierer {verlierer}",
           "Veränderung des Mengenanteils in Prozentpunkten vs. Vorjahr, je Marke in ihrer Kategorie · ★ = eigene Marke")
    cd = CategoryChartData()
    cd.categories = [z[0] for z in zeilen]
    cd.add_series("Δ pp vs. VJ", [round(z[1], 2) for z in zeilen])
    gf = f.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, RAND, Inches(1.8), BREITE - 2 * RAND, Inches(5.0), cd)
    gf.name = "Gewinner und Verlierer"
    chart = gf.chart
    _diagramm_stil(chart, legende=False)
    _achsen(chart, zahlformat='+0.0;-0.0;0')
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '+0.0" pp";-0.0" pp"'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(10)
    serie = chart.series[0]
    serie.invert_if_negative = False
    for i, (name, wert, eigen) in enumerate(zeilen):
        punkt = serie.points[i]
        punkt.format.fill.solid()
        punkt.format.fill.fore_color.rgb = rgb(config.DIVERGIEREND_PLUS if wert >= 0 else config.DIVERGIEREND_MINUS)
    _fusszeile(f, fuss, seite)
    return f, zeilen


def _segment_folie(prs, kat: str, marke: str, matrix: pd.DataFrame, beitraege: pd.DataFrame, fuss, seite):
    f = prs.slides.add_slide(prs.slide_layouts[6])
    gesamt = beitraege["beitrag"].sum()
    beitraege = beitraege.sort_values("beitrag", ascending=gesamt < 0)  # wichtigster Beitrag zuerst
    haupt = beitraege.iloc[0]
    wort = "Verlust" if gesamt < 0 else "Zugewinn"
    if gesamt and haupt["beitrag"] / gesamt >= 0.4:
        aussage = f"{wort} konzentriert sich auf {haupt['kanal']} {haupt['region'].replace('Sued', 'Süd')}"
    else:
        aussage = f"{wort} breit über alle Regionen und Kanäle"
    _titel(f, f"{marke}: {aussage}",
           f"{kat} · Anteilsveränderung von {marke} in pp vs. Vorjahr je Region und Kanal")
    regionen = sorted(matrix["region"].unique(), key=lambda r: ["Nord", "West", "Sued", "Ost"].index(r)
                      if r in ["Nord", "West", "Sued", "Ost"] else 9)
    kanaele = list(dict.fromkeys(matrix["kanal"]))
    zeilen, spalten = len(regionen) + 1, len(kanaele) + 1
    tab = f.shapes.add_table(zeilen, spalten, RAND, Inches(1.9), Inches(7.4), Inches(0.62) * zeilen).table
    tab.columns[0].width = Inches(1.4)
    for j in range(1, spalten):
        tab.columns[j].width = int((Inches(7.4) - Inches(1.4)) / (spalten - 1))

    def zelle(i, j, text, fuell, fett=False, farbe=config.TINTE, groesse=14):
        c = tab.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(fuell)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame
        for k, zeile in enumerate(text.split("\n")):  # jede Zeile ein eigener, zentrierter Absatz
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = zeile
            r.font.size, r.font.name = Pt(groesse if k == 0 else groesse - 2), SCHRIFT
            r.font.bold = fett and k == 0
            r.font.color.rgb = rgb(farbe)

    zelle(0, 0, "", "FFFFFF")
    for j, k in enumerate(kanaele, 1):
        zelle(0, j, k, config.DUNKEL, True, "FFFFFF", 13)
    for i, reg in enumerate(regionen, 1):
        zelle(i, 0, reg.replace("Sued", "Süd"), "F3F6FB", True, config.TINTE, 13)
        for j, k in enumerate(kanaele, 1):
            z = matrix[(matrix["region"] == reg) & (matrix["kanal"] == k)]
            if z.empty:
                zelle(i, j, "–", "FFFFFF")
                continue
            d = float(z["delta_vj"].iloc[0])
            fuell = config.divergierende_farbe(d, 3.0)
            dunkel = abs(d) >= 2.5
            zelle(i, j, f"{pp(d)}\n{zahl(float(z['anteil'].iloc[0]))} %", fuell, abs(d) >= 1,
                  "FFFFFF" if dunkel else config.TINTE, 13)
    _text(f, RAND, Inches(1.9) + Inches(0.62) * zeilen + Inches(0.15), Inches(7.4), Inches(0.4),
          "Zelle: Veränderung vs. VJ und aktueller Anteil · blau = Zugewinn, rot = Verlust", groesse=11,
          farbe=config.TINTE_GEDAEMPFT)
    # Beitraege rechts
    _karte(f, Inches(8.4), Inches(1.9), BREITE - Inches(8.4) - RAND, Inches(4.6), "F7F7F5")
    _text(f, Inches(8.7), Inches(2.1), Inches(4), Inches(0.4), "Beitrag zur nationalen Veränderung", groesse=16,
          fett=True, farbe=config.DUNKEL)
    punkte = [f"National: {pp(gesamt)} vs. VJ"]
    for _, b in beitraege.head(3).iterrows():
        punkte.append(f"{b['kanal']} {b['region'].replace('Sued', 'Süd')}: {pp(b['beitrag'])} "
                      f"(Segmentanteil {pp(b['delta_vj'])})")
    punkte.append("Beitrag = Anteilsveränderung im Segment × Gewicht des Segments (plus Mix-Effekt)")
    _aufzaehlung(f, Inches(8.7), Inches(2.6), BREITE - Inches(8.7) - RAND - Inches(0.3), Inches(3.8), punkte, groesse=13)
    _fusszeile(f, fuss, seite)
    return f


def _prognose_folie(prs, p, fuss, seite):
    f = prs.slides.add_slide(prs.slide_layouts[6])
    ende = p.tabelle.iloc[-1]
    ist = p.historie["anteil"].iloc[-1]
    richtung = "stabil" if abs(ende["prognose"] - ist) < 0.3 else ("steigend" if ende["prognose"] > ist else "sinkend")
    _titel(f, f"Prognose {p.wer}: {richtung} bei {zahl(ende['prognose'])} % bis {ende['monat']}",
           f"{p.kategorie} · Mengenanteil in %, Ist (24 Monate) und Prognose mit 80-%-Band")
    hist = p.historie["anteil"].tail(24)
    monate = list(hist.index) + list(p.tabelle["monat"])
    n_h = len(hist)
    cd = CategoryChartData()
    cd.categories = [m[2:].replace("-", "/") for m in monate]
    cd.add_series("Ist", [round(v, 2) for v in hist] + [None] * len(p.tabelle))
    # Prognose beginnt am letzten Ist-Punkt, damit die Linie anschliesst
    cd.add_series("Prognose", [None] * (n_h - 1) + [round(ist, 2)] + [round(v, 2) for v in p.tabelle["prognose"]])
    cd.add_series("Band oben", [None] * n_h + [round(v, 2) for v in p.tabelle["oben"]])
    cd.add_series("Band unten", [None] * n_h + [round(v, 2) for v in p.tabelle["unten"]])
    gf = f.shapes.add_chart(XL_CHART_TYPE.LINE, RAND, Inches(1.8), Inches(8.0), Inches(5.0), cd)
    gf.name = "Prognose"
    chart = gf.chart
    _diagramm_stil(chart)
    alle = list(hist) + list(p.tabelle["unten"]) + list(p.tabelle["oben"])
    _achsen(chart, min_=max(0, int(min(alle)) - 1), max_=int(max(alle)) + 2)
    stil = [(config.FARBE_EIGEN, 2.75, None), (config.FARBE_EIGEN, 2.75, MSO_LINE_DASH_STYLE.DASH),
            ("86B6EF", 1.25, MSO_LINE_DASH_STYLE.ROUND_DOT), ("86B6EF", 1.25, MSO_LINE_DASH_STYLE.ROUND_DOT)]
    for serie, (farbe, breite, strich) in zip(chart.series, stil):
        serie.smooth = False
        serie.format.line.color.rgb = rgb(farbe)
        serie.format.line.width = Pt(breite)
        serie.marker.style = XL_MARKER_STYLE.NONE
        if strich:
            serie.format.line.dash_style = strich
    # Kennzahlen rechts
    x = Inches(8.9)
    w = BREITE - x - RAND
    _karte(f, x, Inches(1.8), w, Inches(1.5))
    rb = p.rueckblick
    genau = f"± {zahl(rb['mae_kombi'].iloc[min(2, len(rb) - 1)])} pp" if len(rb) else "–"
    _text(f, x + Inches(0.25), Inches(1.95), w - Inches(0.5), Inches(0.35), "Treffsicherheit (3 Monate voraus)",
          groesse=13, farbe=config.TINTE_2)
    _text(f, x + Inches(0.25), Inches(2.3), w - Inches(0.5), Inches(0.7), genau, groesse=34, fett=True,
          farbe=config.DUNKEL)
    _text(f, x + Inches(0.25), Inches(2.95), w - Inches(0.5), Inches(0.3), "mittlerer Fehler im Rückblicktest",
          groesse=11, farbe=config.TINTE_GEDAEMPFT)
    punkte = []
    for name, t in p.szenarien.items():
        punkte.append(f"Szenario {name}: {zahl(t['prognose'].iloc[-1])} % ({pp(t['prognose'].iloc[-1] - ende['prognose'])})")
    for name, wert in p.sensitivitaet.items():
        punkte.append(f"{name} → {pp(wert)}")
    if p.hinweis:
        punkte.append(p.hinweis)
    _text(f, x, Inches(3.55), w, Inches(0.4), "Szenarien und Hebel", groesse=16, fett=True, farbe=config.DUNKEL)
    _aufzaehlung(f, x, Inches(4.0), w, Inches(2.8), punkte or ["Keine Szenarien definiert"], groesse=13)
    _fusszeile(f, fuss, seite)
    f.notes_slide.notes_text_frame.text = (
        f"Kombination aus Treibermodell (Gewicht {zahl(p.gewicht_treiber * 100, 0)} %) und Trendfortschreibung. "
        f"Planannahme: Preis, Distribution und Werbedruck auf dem Niveau der letzten 3 Monate.")
    return f


def _methodik_folie(prs, pruefungen, quellen, fuss, seite):
    f = prs.slides.add_slide(prs.slide_layouts[6])
    _titel(f, "Datenbasis und Methodik")
    links = [f"Quellen: {', '.join(q['datei'] for q in quellen)}"]
    links += [f"{p.stufe}: {p.text}" for p in pruefungen if p.stufe != "OK"][:5] or ["Datenqualität: keine Auffälligkeiten"]
    links += [p.text for p in pruefungen if p.stufe == "OK"][:2]
    _text(f, RAND, Inches(1.6), Inches(5.8), Inches(0.4), "Datenqualität (automatisch geprüft)", groesse=16, fett=True,
          farbe=config.DUNKEL)
    _aufzaehlung(f, RAND, Inches(2.1), Inches(5.8), Inches(4.5), [kuerze(t, 160) for t in links], groesse=13)
    rechts = [
        "Mengenanteil = Absatz Marke / Absatz Kategorie; Veränderungen in Prozentpunkten (pp)",
        "Standard: rollierende 3 Monate vs. gleicher Zeitraum Vorjahr (VJ) und vorherige 3 Monate (VP)",
        "Preisindex = Durchschnittspreis Marke / Durchschnittspreis Kategorie × 100",
        "Beitragsanalyse zerlegt die nationale Veränderung exakt in Segmentbeiträge (Performance + Mix)",
        "Prognose: Treibermodell (Preis, Distribution, Werbedruck) kombiniert mit gedämpftem Trend; "
        "Gewichte und 80-%-Band aus Rückblicktest",
        "Treiberwirkungen und Szenarien sind statistische Schätzungen aus der Historie, keine Kausalbeweise",
    ]
    _text(f, Inches(7.0), Inches(1.6), Inches(5.7), Inches(0.4), "Definitionen", groesse=16, fett=True,
          farbe=config.DUNKEL)
    _aufzaehlung(f, Inches(7.0), Inches(2.1), BREITE - Inches(7.0) - RAND, Inches(4.5), rechts, groesse=13)
    _fusszeile(f, fuss, seite)
    return f


def schreibe(pfad: Path, titel: str, zeitraum_name: str, ueberblick: pd.DataFrame, rankings: dict,
             zeitreihen: dict, eigene: dict, erkenntnisse: list[Erkenntnis], segment: tuple | None,
             prognosen: dict, pruefungen, quellen, hinweis: str = "") -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = BREITE, HOEHE
    prs.core_properties.title = titel
    prs.core_properties.author = "Insights-Automation"
    fuss = f"{titel} · {zeitraum_name}"
    _titelfolie(prs, titel, f"Zeitraum {zeitraum_name} · Vergleich zum Vorjahr",
                hinweis or f"Automatisch erstellt am {date.today().strftime('%d.%m.%Y')}")
    seite = 2
    kern = [kuerze(e.text, 200) for e in erkenntnisse if e.typ in ("Marke", "Neu", "Auffaellig")][:4]
    _zusammenfassung(prs, ueberblick, kern or [e.text for e in erkenntnisse[:4]], fuss, seite)
    for kat, zr in zeitreihen.items():
        seite += 1
        _kategorie_folie(prs, kat, zr, eigene.get(kat, []), pro_kategorie(erkenntnisse, kat, max_anzahl=8), fuss, seite)
    seite += 1
    _gewinner_verlierer(prs, rankings, fuss, seite)
    if segment:
        seite += 1
        _segment_folie(prs, *segment, fuss, seite)
    for p in prognosen.values():
        seite += 1
        _prognose_folie(prs, p, fuss, seite)
    seite += 1
    _methodik_folie(prs, pruefungen, quellen, fuss, seite)
    pfad.parent.mkdir(parents=True, exist_ok=True)
    prs.save(pfad)
    return pfad

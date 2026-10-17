"""Excel-Report mit formatierten Tabellen und echten (bearbeitbaren) Excel-Diagrammen."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .kommentar import Erkenntnis
from .laden import Pruefung

KOPF = PatternFill("solid", fgColor=config.DUNKEL)
KOPF_FONT = Font(bold=True, color="FFFFFF")
TITEL_FONT = Font(bold=True, size=14, color=config.DUNKEL)
LINIE = Border(bottom=Side(style="thin", color=config.GITTER))


def _tabelle(ws, df: pd.DataFrame, start_zeile: int, formate: dict[str, str] | None = None,
             delta_spalten: tuple[str, ...] = (), delta_grenze: float = 1.0, start_spalte: int = 1) -> int:
    formate = formate or {}
    for j, spalte in enumerate(df.columns, start_spalte):
        c = ws.cell(row=start_zeile, column=j, value=spalte)
        c.fill, c.font = KOPF, KOPF_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (_, zeile) in enumerate(df.iterrows(), start_zeile + 1):
        for j, spalte in enumerate(df.columns, start_spalte):
            wert = zeile[spalte]
            if isinstance(wert, float) and pd.isna(wert):
                wert = None
            c = ws.cell(row=i, column=j, value=wert)
            c.border = LINIE
            if spalte in formate:
                c.number_format = formate[spalte]
            if spalte in delta_spalten and isinstance(wert, (int, float)):
                c.fill = PatternFill("solid", fgColor=config.divergierende_farbe(wert, delta_grenze))
    for j, spalte in enumerate(df.columns, start_spalte):
        breite = max([len(str(spalte))] + [len(str(v)) for v in df[spalte].head(50)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, breite + 2), 60)
    ws.row_dimensions[start_zeile].height = 32
    return start_zeile + len(df) + 2


PP = '+0.0" pp";-0.0" pp";0.0" pp"'
PROZ = '0.0" %"'
PROZ_DELTA = '+0.0" %";-0.0" %";0.0" %"'
GANZ = '#,##0'


def schreibe(pfad: Path, titel: str, zeitraum_name: str, ueberblick: pd.DataFrame,
             rankings: dict[str, pd.DataFrame], segmente: dict[str, tuple[str, pd.DataFrame]],
             zeitreihen: dict[str, pd.DataFrame], prognosen: dict, erkenntnisse: list[Erkenntnis],
             pruefungen: list[Pruefung]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Übersicht"
    ws["A1"], ws["A1"].font = titel, TITEL_FONT
    ws["A2"] = f"Zeitraum: {zeitraum_name}  ·  Veränderungen vs. Vorjahr (VJ) und Vorperiode (VP)  ·  Mengenanteile"
    ws["A2"].font = Font(color=config.TINTE_2)
    u = ueberblick.rename(columns={
        "kategorie": "Kategorie", "markt_absatz": "Markt Absatz", "markt_veraenderung_vj": "Markt vs. VJ",
        "anteil_menge": "Anteil Menge", "delta_menge_vj": "Δ VJ", "delta_menge_vp": "Δ VP",
        "anteil_wert": "Anteil Wert", "delta_wert_vj": "Δ Wert VJ", "preisindex": "Preisindex",
        "distribution": "Distribution gew.", "rang": "Rang"})
    z = _tabelle(ws, u, 4, {"Markt Absatz": GANZ, "Markt vs. VJ": PROZ_DELTA, "Anteil Menge": PROZ, "Δ VJ": PP,
                            "Δ VP": PP, "Anteil Wert": PROZ, "Δ Wert VJ": PP, "Preisindex": "0",
                            "Distribution gew.": PROZ}, delta_spalten=("Δ VJ", "Δ VP", "Δ Wert VJ"))
    ws.cell(row=z, column=1, value="Wichtigste Erkenntnisse").font = Font(bold=True, size=12, color=config.DUNKEL)
    for i, e in enumerate(erkenntnisse[:12], z + 1):
        k = ws.cell(row=i, column=1, value=e.kategorie)
        k.font = Font(color=config.TINTE_2)
        k.alignment = Alignment(vertical="top")
        c = ws.cell(row=i, column=2, value=e.text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=11)
        ws.row_dimensions[i].height = 32 if len(e.text) < 140 else 48
    ws.column_dimensions["A"].width = 16

    ws = wb.create_sheet("Erkenntnisse")
    alle = pd.DataFrame([{"Kategorie": e.kategorie, "Typ": e.typ, "Eigene": "ja" if e.eigen else "",
                          "Erkenntnis": e.text} for e in erkenntnisse])
    _tabelle(ws, alle, 1)
    ws.column_dimensions["D"].width = 120
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    for kat, r in rankings.items():
        ws = wb.create_sheet(f"Ranking {kat}"[:31])
        ws["A1"], ws["A1"].font = f"Markenranking {kat} – {zeitraum_name}", TITEL_FONT
        t = r[["hersteller", "marke", "absatz", "anteil_menge", "delta_menge_vj", "delta_menge_vp", "anteil_wert",
               "preisindex", "delta_preisindex_vj", "distribution", "delta_distribution_vj"]].rename(columns={
            "hersteller": "Hersteller", "marke": "Marke", "absatz": "Absatz", "anteil_menge": "Anteil Menge",
            "delta_menge_vj": "Δ VJ", "delta_menge_vp": "Δ VP", "anteil_wert": "Anteil Wert",
            "preisindex": "Preisindex", "delta_preisindex_vj": "Δ Preisindex VJ", "distribution": "Distribution gew.",
            "delta_distribution_vj": "Δ Distribution VJ"})
        _tabelle(ws, t, 3, {"Absatz": GANZ, "Anteil Menge": PROZ, "Δ VJ": PP, "Δ VP": PP, "Anteil Wert": PROZ,
                            "Preisindex": "0", "Δ Preisindex VJ": "+0;-0;0", "Distribution gew.": PROZ,
                            "Δ Distribution VJ": "+0.0;-0.0;0.0"}, delta_spalten=("Δ VJ", "Δ VP"))
        for zeile in range(4, 4 + len(t)):
            if ws.cell(row=zeile, column=1).value == config.EIGENER_HERSTELLER:
                for spalte in range(1, 3):
                    ws.cell(row=zeile, column=spalte).font = Font(bold=True, color=config.FARBE_EIGEN)

    ws = wb.create_sheet("Segmente")
    zeile = 1
    for kat, (marke, m) in segmente.items():
        ws.cell(row=zeile, column=1, value=f"{kat}: {marke} – Anteil und Δ vs. VJ je Region × Kanal").font = TITEL_FONT
        piv = m.pivot(index="region", columns="kanal", values="delta_vj").round(2)
        anteil = m.pivot(index="region", columns="kanal", values="anteil").round(1)
        t = piv.reset_index().rename(columns={"region": "Region (Δ pp)"})
        t["Region (Δ pp)"] = t["Region (Δ pp)"].str.replace("Sued", "Süd")
        nach = _tabelle(ws, t, zeile + 1, {c: PP for c in piv.columns}, delta_spalten=tuple(piv.columns),
                        delta_grenze=3.0)
        t2 = anteil.reset_index().rename(columns={"region": "Region (Anteil %)"})
        t2["Region (Anteil %)"] = t2["Region (Anteil %)"].str.replace("Sued", "Süd")
        zeile = _tabelle(ws, t2, nach - 1, {c: PROZ for c in anteil.columns}) + 1

    ws = wb.create_sheet("Zeitreihen")
    spalte = 1
    n_kat = 0
    for kat, piv in zeitreihen.items():
        ws.cell(row=1, column=spalte, value=f"Mengenanteile {kat} (%)").font = TITEL_FONT
        t = piv.reset_index()
        _tabelle(ws, t, 2, {c: '0.0' for c in piv.columns}, start_spalte=spalte)
        chart = LineChart()
        chart.title = f"Mengenanteile {kat}"
        chart.y_axis.title = "Anteil %"
        chart.height, chart.width = 8, 18
        daten = Reference(ws, min_col=spalte + 1, max_col=spalte + len(piv.columns), min_row=2, max_row=2 + len(t))
        kat_ref = Reference(ws, min_col=spalte, min_row=3, max_row=2 + len(t))
        chart.add_data(daten, titles_from_data=True)
        chart.set_categories(kat_ref)
        farben = iter(config.FARBEN_WETTBEWERB)
        for serie, name in zip(chart.series, piv.columns):
            eigen = name in set(rankings[kat].query("hersteller == @config.EIGENER_HERSTELLER")["marke"])
            farbe = config.FARBE_EIGEN if eigen else next(farben, config.FARBE_UEBRIGE)
            serie.graphicalProperties.line.solidFill = farbe
            serie.graphicalProperties.line.width = 28000 if eigen else 19000
            serie.smooth = False
        # Diagramme untereinander unter den Tabellen, damit sie sich nicht ueberlappen
        ws.add_chart(chart, f"A{len(t) + 5 + n_kat * 18}")
        n_kat += 1
        # Ueber die Tabellen hinweg Platz lassen
        for c in range(spalte, spalte + len(piv.columns) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 11
        spalte += len(piv.columns) + 2

    if prognosen:
        ws = wb.create_sheet("Prognose")
        zeile = 1
        for (kat, wer), p in prognosen.items():
            ws.cell(row=zeile, column=1, value=f"Prognose Mengenanteil {wer} ({kat})").font = TITEL_FONT
            hist = p.historie[["anteil"]].tail(24).reset_index().rename(columns={"anteil": "Ist"})
            prog = p.tabelle.rename(columns={"prognose": "Prognose", "unten": "Band unten", "oben": "Band oben"})
            t = pd.concat([hist, prog[["monat", "Prognose", "Band unten", "Band oben"]]], ignore_index=True)
            t = t.rename(columns={"monat": "Monat"})
            ende = _tabelle(ws, t, zeile + 1, {c: '0.00' for c in t.columns[1:]})
            chart = LineChart()
            chart.title = f"{wer}: Ist und Prognose (80-%-Band)"
            chart.height, chart.width = 8, 18
            daten = Reference(ws, min_col=2, max_col=5, min_row=zeile + 1, max_row=zeile + 1 + len(t))
            chart.add_data(daten, titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=zeile + 2, max_row=zeile + 1 + len(t)))
            stile = [(config.FARBE_EIGEN, 28000, None), (config.FARBE_EIGEN, 28000, "dash"),
                     (config.FARBE_PROGNOSE_BAND, 15000, None), (config.FARBE_PROGNOSE_BAND, 15000, None)]
            for serie, (farbe, breite, strich) in zip(chart.series, stile):
                serie.graphicalProperties.line.solidFill = farbe
                serie.graphicalProperties.line.width = breite
                if strich:
                    serie.graphicalProperties.line.dashStyle = strich
            ws.add_chart(chart, f"G{zeile + 1}")
            for c in "GHIJ":
                ws.column_dimensions[c].width = 15
            # Rueckblicktest und Szenarien rechts unter dem Diagramm
            info = zeile + 19
            ws.cell(row=info, column=7, value="Rückblicktest: mittlerer Fehler (pp) je Prognosehorizont").font = Font(bold=True)
            if len(p.rueckblick):
                rb = p.rueckblick.rename(columns={"horizont": "Monate voraus", "mae_treiber": "Treibermodell",
                                                  "mae_basis": "Trendmodell", "mae_kombi": "Kombination"})
                rb = rb[["Monate voraus", "Treibermodell", "Trendmodell", "Kombination"]]
                for j, n in enumerate(rb.columns):
                    c = ws.cell(row=info + 1, column=7 + j, value=n)
                    c.fill, c.font = KOPF, KOPF_FONT
                for i, (_, row) in enumerate(rb.iterrows()):
                    for j, v in enumerate(row):
                        c = ws.cell(row=info + 2 + i, column=7 + j, value=float(v))
                        c.number_format = "0" if j == 0 else "0.00"
                info += len(rb) + 3
            if p.szenarien:
                ws.cell(row=info, column=7, value="Szenarien (Anteil am Ende des Horizonts)").font = Font(bold=True)
                ws.cell(row=info + 1, column=7, value="Basis")
                ws.cell(row=info + 1, column=8, value=float(p.tabelle["prognose"].iloc[-1])).number_format = "0.00"
                for i, (name, t_s) in enumerate(p.szenarien.items(), 2):
                    ws.cell(row=info + i, column=7, value=name)
                    ws.cell(row=info + i, column=8, value=float(t_s["prognose"].iloc[-1])).number_format = "0.00"
            zeile = max(ende, info + 8) + 2

    ws = wb.create_sheet("Datenqualität")
    _tabelle(ws, pd.DataFrame([{"Stufe": p.stufe, "Thema": p.thema, "Befund": p.text} for p in pruefungen]), 1)
    ws.column_dimensions["C"].width = 100

    for blatt in wb.worksheets:  # beim Drucken: Querformat, eine Seite breit
        blatt.page_setup.orientation = "landscape"
        blatt.sheet_properties.pageSetUpPr.fitToPage = True
        blatt.page_setup.fitToWidth, blatt.page_setup.fitToHeight = 1, 0
        blatt.sheet_view.showGridLines = False
    pfad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pfad)
    return pfad
